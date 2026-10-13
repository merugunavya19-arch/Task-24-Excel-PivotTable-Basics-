from pathlib import Path
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter

ROOT=Path('/home/ubuntu/Sample-Superstore-PivotTables')
ROOT.mkdir(parents=True, exist_ok=True)
CSV=Path('/home/ubuntu/work/superstore.csv')
OUT=ROOT/'Excel-PivotTable-Basics.xlsx'

df=pd.read_csv(CSV)
df['Order Date']=pd.to_datetime(df['Order Date'])
df['Ship Date']=pd.to_datetime(df['Ship Date'])
df['Month']=df['Order Date'].dt.to_period('M').astype(str)

cat=(df.groupby('Category',as_index=False).agg(
    Sales=('Sales','sum'), Profit=('Profit','sum'), Quantity=('Quantity','sum'),
    Orders=('Order ID','nunique'), Avg_Discount=('Discount','mean'))
    .sort_values('Sales',ascending=False))
reg=(df.groupby('Region',as_index=False).agg(
    Sales=('Sales','sum'), Profit=('Profit','sum'), Quantity=('Quantity','sum'),
    Orders=('Order ID','nunique'), Avg_Discount=('Discount','mean'))
    .sort_values('Sales',ascending=False))
month=(df.groupby('Month',as_index=False).agg(
    Sales=('Sales','sum'), Profit=('Profit','sum'), Quantity=('Quantity','sum'),
    Orders=('Order ID','nunique'))
    .sort_values('Month'))

# Keep original source columns in the raw data sheet, plus Month for easy filtering.
raw=df.copy()
raw['Order Date']=raw['Order Date'].dt.date
raw['Ship Date']=raw['Ship Date'].dt.date

wb=Workbook()
ws=wb.active; ws.title='Read Me'
for name in ['Superstore_Data','Pivot_Category','Pivot_Region','Pivot_Month','Insights']:
    wb.create_sheet(name)

navy='17365D'; blue='2F75B5'; light='D9EAF7'; pale='F4F8FB'; green='E2F0D9'; orange='FCE4D6'; white='FFFFFF'; gray='666666'
thin=Side(style='thin',color='D9E2F3')

def title(ws, text, subtitle=None):
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=8)
    c=ws.cell(1,1,text); c.font=Font(size=16,bold=True,color=white); c.fill=PatternFill('solid',fgColor=navy); c.alignment=Alignment(horizontal='left')
    ws.row_dimensions[1].height=28
    if subtitle:
        ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=8)
        c=ws.cell(2,1,subtitle); c.font=Font(size=10,italic=True,color=gray); c.alignment=Alignment(wrap_text=True)

def style_header(ws,row,cols):
    for col in range(1,cols+1):
        c=ws.cell(row,col); c.font=Font(bold=True,color=white); c.fill=PatternFill('solid',fgColor=blue); c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        c.border=Border(bottom=thin)
    ws.row_dimensions[row].height=24

def format_table(ws, start_row, start_col, end_row, end_col, name):
    ref=f'{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{end_row}'
    tab=Table(displayName=name, ref=ref)
    tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showFirstColumn=False,showLastColumn=False,showRowStripes=True,showColumnStripes=False)
    ws.add_table(tab)

def widths(ws, max_col):
    for col in range(1,max_col+1):
        letter=get_column_letter(col); maxlen=0
        for cell in ws[letter]:
            if cell.value is not None: maxlen=max(maxlen,len(str(cell.value)))
        ws.column_dimensions[letter].width=min(max(maxlen+2,12),34)

# Read Me
r=ws
r.merge_cells('A1:H1'); r['A1']='Excel PivotTable Basics'; r['A1'].font=Font(size=18,bold=True,color=white); r['A1'].fill=PatternFill('solid',fgColor=navy); r['A1'].alignment=Alignment(horizontal='left'); r.row_dimensions[1].height=32
r.merge_cells('A2:H2'); r['A2']='Sample Superstore | Sales summaries by category, region, and month'; r['A2'].font=Font(size=11,italic=True,color=gray)
info=[('Repository name','Excel-PivotTable-Basics'),('Task','Create PivotTables for sales by category, region, and month'),('Dataset','Sample Superstore sales dataset'),('Source','Public dataset: github.com/leonism/sample-superstore'),('Records',len(df)),('Workbook sheets','Raw data + 3 pivot-style summaries + insights')]
for i,(k,v) in enumerate(info,4):
 r.cell(i,1,k).font=Font(bold=True,color=navy); r.cell(i,2,v)
 r.cell(i,1).fill=PatternFill('solid',fgColor=light); r.cell(i,2).fill=PatternFill('solid',fgColor=pale)
r['A12']='How to use'; r['A12'].font=Font(size=13,bold=True,color=navy)
notes=['Start with Superstore_Data for the cleaned source table.','Review Pivot_Category, Pivot_Region, and Pivot_Month for grouped sales summaries.','Use the Insights sheet for the main findings.','Amounts are shown in currency format; discounts are shown as percentages.']
for j,n in enumerate(notes,13): r.cell(j,1,'• '+n)
r.freeze_panes='A4'; widths(r,8)

# Raw data
rd=wb['Superstore_Data']; title(rd,'Sample Superstore Data','Source records used for the three summary tables. Filterable Excel table with dates, sales, quantity, discount, and profit.')
raw_cols=list(raw.columns)
for j,c in enumerate(raw_cols,1): rd.cell(4,j,c)
style_header(rd,4,len(raw_cols))
for i,row in enumerate(raw.itertuples(index=False,name=None),5):
    for j,v in enumerate(row,1): rd.cell(i,j,v)
format_table(rd,4,1,4+len(raw),len(raw_cols),'SuperstoreData')
rd.freeze_panes='A5'; rd.auto_filter.ref=f'A4:{get_column_letter(len(raw_cols))}{4+len(raw)}'
for col in ['R','S','T','U']:
    if col in rd.column_dimensions: pass
# identify numeric/date columns by header
headers={rd.cell(4,c).value:c for c in range(1,len(raw_cols)+1)}
for row in range(5,5+len(raw)):
    for h in ['Sales','Profit']:
        rd.cell(row,headers[h]).number_format='$#,##0.00'
    rd.cell(row,headers['Discount']).number_format='0.0%'
for h in ['Order Date','Ship Date']:
    for row in range(5,5+len(raw)): rd.cell(row,headers[h]).number_format='m/d/yyyy'
widths(rd,len(raw_cols))

# summary sheets
summary_specs=[('Pivot_Category',cat,'Sales by Category','CategoryPivot'),('Pivot_Region',reg,'Sales by Region','RegionPivot'),('Pivot_Month',month,'Sales by Month','MonthPivot')]
for sheet,data,subtitle,tname in summary_specs:
    sw=wb[sheet]; title(sw,subtitle,'Pivot-style summary: grouped totals generated from the Sample Superstore source data.')
    cols=list(data.columns)
    for j,c in enumerate(cols,1): sw.cell(4,j,c.replace('_',' '))
    style_header(sw,4,len(cols))
    for i,row in enumerate(data.itertuples(index=False,name=None),5):
        for j,v in enumerate(row,1): sw.cell(i,j,v)
    format_table(sw,4,1,4+len(data),len(cols),tname)
    sw.freeze_panes='A5'; widths(sw,len(cols))
    for row in range(5,5+len(data)):
        for h in ['Sales','Profit','AveragePrice','MinimumPrice','MaximumPrice','InventoryValue']:
            if h in cols: sw.cell(row,cols.index(h)+1).number_format='$#,##0.00'
        for h in ['Avg_Discount']:
            if h in cols: sw.cell(row,cols.index(h)+1).number_format='0.0%'
    # Add chart for each pivot-style summary.
    chart=BarChart() if sheet!='Pivot_Month' else LineChart()
    chart.type='col'; chart.style=10; chart.title=subtitle; chart.y_axis.title='Sales'; chart.x_axis.title=cols[0].replace('_',' ')
    data_ref=Reference(sw,min_col=2,min_row=4,max_row=4+len(data))
    cats_ref=Reference(sw,min_col=1,min_row=5,max_row=4+len(data))
    chart.add_data(data_ref,titles_from_data=True); chart.set_categories(cats_ref); chart.height=7; chart.width=13
    sw.add_chart(chart,'H4')

# Insights
ins=wb['Insights']; title(ins,'Key Insights','Short observations from the three pivot-style summaries.')
ins['A4']='Metric'; ins['B4']='Finding'; style_header(ins,4,2)
insights=[
 ('Top category by sales',f"{cat.iloc[0]['Category']} leads sales with ${cat.iloc[0]['Sales']:,.2f}."),
 ('Top region by sales',f"{reg.iloc[0]['Region']} is the highest-sales region with ${reg.iloc[0]['Sales']:,.2f}."),
 ('Most profitable category',f"{cat.sort_values('Profit',ascending=False).iloc[0]['Category']} has the highest total profit."),
 ('Best sales month',f"{month.loc[month['Sales'].idxmax(),'Month']} has the highest monthly sales at ${month['Sales'].max():,.2f}."),
 ('Profitability check',f"Overall profit is ${df['Profit'].sum():,.2f} on ${df['Sales'].sum():,.2f} in sales."),
 ('Business takeaway','Use the category and region summaries to prioritize high-sales areas, then review profit and discount together before making pricing decisions.')]
for i,(m,f) in enumerate(insights,5): ins.cell(i,1,m); ins.cell(i,2,f); ins.cell(i,2).alignment=Alignment(wrap_text=True,vertical='top'); ins.cell(i,1).fill=PatternFill('solid',fgColor=green if i<10 else orange)
ins.column_dimensions['A'].width=26; ins.column_dimensions['B'].width=88; ins.freeze_panes='A5'

# Global workbook polish
for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines=False
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None and cell.row not in [1,2,4]:
                cell.alignment = cell.alignment.copy(vertical='center')
wb.calculation.fullCalcOnLoad=True
wb.calculation.forceFullCalc=True
wb.save(OUT)
print(OUT)
print('records',len(df),'category rows',len(cat),'region rows',len(reg),'month rows',len(month))
